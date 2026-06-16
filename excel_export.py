from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _autosize_and_style(worksheet):
    """تطبيق تنسيق احترافي تلقائي على كل ورقة."""
    worksheet.sheet_view.rightToLeft = True
    worksheet.freeze_panes = "A2"

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'

    for column_cells in worksheet.columns:
        max_length = 12
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, min(len(str(cell.value)) + 2, 45))
        worksheet.column_dimensions[column_letter].width = max_length


def _write_sheet(writer, sheet_name: str, frame: pd.DataFrame):
    safe_frame = frame.copy()
    if safe_frame.empty:
        safe_frame = pd.DataFrame({"البيان": ["لا توجد بيانات"]})
    safe_frame.to_excel(writer, sheet_name=sheet_name, index=False)
    _autosize_and_style(writer.book[sheet_name])


def export_to_excel(
    metrics: dict,
    customers: pd.DataFrame,
    products: pd.DataFrame,
    representatives: pd.DataFrame,
    branches: pd.DataFrame,
    insights: pd.DataFrame,
) -> BytesIO:
    """إنشاء ملف Excel متعدد الأوراق وجاهز للتنزيل."""
    output = BytesIO()

    dashboard = pd.DataFrame(
        [
            {"المؤشر": "إجمالي المبيعات الحالية", "القيمة": metrics["current_total"]},
            {"المؤشر": "إجمالي المبيعات السابقة", "القيمة": metrics["previous_total"]},
            {"المؤشر": "إجمالي الفرق", "القيمة": metrics["difference"]},
            {"المؤشر": "نسبة النمو", "القيمة": metrics["growth"]},
            {"المؤشر": "عدد العملاء", "القيمة": metrics["customers_count"]},
            {"المؤشر": "عدد المنتجات", "القيمة": metrics["products_count"]},
            {"المؤشر": "عدد المناديب", "القيمة": metrics["representatives_count"]},
            {"المؤشر": "عدد الفروع", "القيمة": metrics["branches_count"]},
        ]
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _write_sheet(writer, "Dashboard", dashboard)
        _write_sheet(writer, "Customers", customers)
        _write_sheet(writer, "Products", products)
        _write_sheet(writer, "Representatives", representatives)
        _write_sheet(writer, "Branches", branches)
        _write_sheet(writer, "Insights", insights)

        for worksheet in writer.book.worksheets:
            worksheet.sheet_properties.pageSetUpPr.fitToPage = True
            worksheet.page_setup.fitToWidth = 1

    output.seek(0)
    return output
